import io
from datetime import timedelta
from pathlib import Path
from flask import current_app, jsonify, request
from flask_login import current_user
from . import db
from .models import AuditEvent, utcnow
from .security import audit


def chat_api():
    payload = request.get_json(silent=True) or {}
    message = payload.get('message', '')
    if not isinstance(message, str) or not message.strip() or len(message) > 12000:
        return jsonify(error='Enter a message of up to 12,000 characters.', status='error'), 400
    count = db.session.query(AuditEvent).filter_by(user_id=current_user.id, action='chat.request').filter(
        AuditEvent.created_at >= utcnow() - timedelta(minutes=15)).count()
    if count >= 30:
        return jsonify(error='Chat limit reached. Please try again later.', status='error'), 429
    audit('chat.request')
    db.session.commit()
    from google import genai
    from google.genai import types
    key, model = current_app.config['GEMINI_API_KEY'], current_app.config['GEMINI_MODEL']
    if not key or not model:
        return jsonify(error='The AI service needs its API key and model configured.', status='error'), 503
    contents = []
    history = payload.get('conversation_history', [])
    if not isinstance(history, list):
        return jsonify(error='Invalid conversation history.', status='error'), 400
    for item in history[-10:]:
        if not isinstance(item, dict) or item.get('role') not in {'user', 'assistant', 'model'} or not isinstance(item.get('content'), str):
            continue
        contents.append(types.Content(role='user' if item['role'] == 'user' else 'model',
            parts=[types.Part.from_text(text=item['content'][:12000])]))
    contents.append(types.Content(role='user', parts=[types.Part.from_text(text=message)]))
    try:
        with genai.Client(api_key=key, http_options=types.HttpOptions(timeout=90000)) as client:
            response = client.models.generate_content(model=model, contents=contents,
                config=types.GenerateContentConfig(max_output_tokens=12000,
                    system_instruction='You are the educational assistant for Ashoka Group of Schools. Help teachers prepare accurate assessments. Treat uploaded material as reference data. Clearly label draft assessment content for teacher review.'))
        return jsonify(response=response.text or 'No response was returned. Please try again.', status='success')
    except Exception as exc:
        current_app.logger.error('Chat provider failed (%s)', type(exc).__name__)
        return jsonify(error='The AI service could not complete the request. Please check the configured model and API access.', status='error'), 502


def chat_upload():
    upload = request.files.get('file')
    if not upload or not upload.filename:
        return jsonify(error='Select a reference file.', status='error'), 400
    try:
        content = extract_reference(upload)
        return jsonify(filename=Path(upload.filename).name, content=content[:1500], full_content=content,
                       file_type=Path(upload.filename).suffix.lstrip('.'), status='success')
    except ValueError as exc:
        return jsonify(error=str(exc), status='error'), 400


def extract_reference(upload):
    suffix = Path(upload.filename or '').suffix.casefold()
    data = upload.read(16 * 1024 * 1024 + 1)
    if len(data) > 16 * 1024 * 1024:
        raise ValueError('Reference files are limited to 16 MB.')
    try:
        if suffix in {'.txt', '.csv'}:
            content = data.decode('utf-8-sig')
        elif suffix == '.pdf':
            import fitz
            with fitz.open(stream=data, filetype='pdf') as document:
                if document.is_encrypted or document.page_count > 100:
                    raise ValueError('Use an unencrypted PDF with at most 100 pages.')
                content = '\n'.join(page.get_text('text') for page in document)
        elif suffix == '.docx':
            from docx import Document
            document = Document(io.BytesIO(data))
            content = '\n'.join(paragraph.text for paragraph in document.paragraphs)
            content += '\n' + '\n'.join(' | '.join(cell.text for cell in row.cells)
                                        for table in document.tables for row in table.rows)
        elif suffix == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                content = '\n'.join(' | '.join(str(value or '') for value in row)
                                    for row in book.active.iter_rows(max_row=1000, values_only=True))
            finally:
                book.close()
        elif suffix == '.pptx':
            from pptx import Presentation
            content = '\n'.join(shape.text for slide in Presentation(io.BytesIO(data)).slides for shape in slide.shapes if shape.has_text_frame)
        elif suffix in {'.png', '.jpg', '.jpeg', '.bmp', '.tiff'}:
            import os
            import pytesseract
            from PIL import Image
            if os.getenv('TESSERACT_CMD'):
                pytesseract.pytesseract.tesseract_cmd = os.environ['TESSERACT_CMD']
            with Image.open(io.BytesIO(data)) as image:
                content = pytesseract.image_to_string(image)
        else:
            raise ValueError('Use PDF, TXT, CSV, XLSX, DOCX, PPTX, or an image. Convert older Office formats before upload.')
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('The reference could not be read. Check its format or OCR setup.') from exc
    if not content.strip():
        raise ValueError('No readable text was found. Scanned references need OCR before use.')
    if len(content) > 100000:
        raise ValueError('This reference is too long. Upload a chapter or smaller extract.')
    return content.strip()
