"""Bounded document and reviewed-bank ingestion; never executes uploaded content."""
import io
import json
import zipfile
from pathlib import Path

from .validation import validate_question


def source_passages(upload):
    data = upload.read(32 * 1024 * 1024 + 1)
    if len(data) > 32 * 1024 * 1024:
        raise ValueError('Source is too large (32 MB maximum).')
    suffix = Path(upload.filename or '').suffix.casefold()
    if suffix == '.pdf':
        import fitz
        if not data.startswith(b'%PDF-'):
            raise ValueError('The uploaded file is not a valid PDF.')
        try:
            with fitz.open(stream=data, filetype='pdf') as document:
                if document.is_encrypted or document.page_count > 60:
                    raise ValueError('Use an unencrypted PDF with at most 60 pages.')
                passages = {}
                for index, page in enumerate(document):
                    value = page.get_text('text').strip()
                    if len(value) < 20:
                        raise ValueError(f'Page {index + 1} has insufficient readable text. Review OCR externally before publishing this source.')
                    passages[f'page-{index + 1}'] = value
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('PDF could not be read.') from exc
    elif suffix == '.txt':
        try:
            value = data.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('Save text sources as UTF-8.') from exc
        passages = {f'passage-{i + 1}': part.strip() for i, part in enumerate(value.split('\n\n')) if part.strip()}
    else:
        raise ValueError('Source uploads support PDF and UTF-8 TXT files.')
    if not passages or len(json.dumps(passages)) > 100000:
        raise ValueError('Source must contain readable text and fit within 100,000 characters. Split large sources by chapter.')
    return passages


BANK_COLUMNS = ['Question', 'Type', 'Marks', 'Chapter', 'Difficulty', 'Bloom', 'Options', 'Answer',
                'Solution', 'Rubric', 'Source Reference', 'Source Quote']


def bank_questions(upload):
    data = upload.read(8 * 1024 * 1024 + 1)
    if len(data) > 8 * 1024 * 1024:
        raise ValueError('Bank imports are limited to 8 MB.')
    suffix = Path(upload.filename or '').suffix.casefold()
    if suffix == '.json':
        try:
            items = json.loads(data.decode('utf-8-sig'))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError('Provide a valid UTF-8 JSON question array.') from exc
    elif suffix == '.xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError('Expanded spreadsheet exceeds the import limit.')
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.active
                iterator = sheet.iter_rows(values_only=True)
                headers = [str(v or '').strip() for v in next(iterator)]
                if not set(BANK_COLUMNS) <= set(headers):
                    raise ValueError('Use the reviewed-bank spreadsheet template; answer and rubric columns are required.')
                items = []
                for row_index, values in enumerate(iterator, start=2):
                    if row_index > 502:
                        raise ValueError('Import at most 500 questions at once.')
                    if not any(v is not None for v in values):
                        continue
                    row = dict(zip(headers, values))
                    if any(isinstance(v, str) and v.startswith('=') for v in values):
                        raise ValueError('Spreadsheet formulas are not accepted in bank imports.')
                    items.append({'stem': str(row['Question'] or ''), 'type': row['Type'],
                        'marks': row['Marks'], 'chapter': row['Chapter'], 'difficulty': row['Difficulty'],
                        'bloom': row['Bloom'], 'options': json.loads(row['Options'] or '[]'),
                        'answer': str(row['Answer'] or ''), 'solution': str(row['Solution'] or ''),
                        'rubric': json.loads(row['Rubric'] or '[]'),
                        'source_reference': str(row['Source Reference'] or ''), 'source_quote': str(row['Source Quote'] or '')})
            finally:
                workbook.close()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('Spreadsheet could not be read. Use the reviewed-bank template.') from exc
    else:
        raise ValueError('Reviewed bank imports support XLSX and JSON files.')
    if not isinstance(items, list) or not 1 <= len(items) <= 500:
        raise ValueError('Import 1–500 reviewed questions at once.')
    return [validate_question(q) for q in items]


def bank_template():
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = 'Reviewed questions'
    sheet.append(BANK_COLUMNS)
    sheet.append(['What is one half plus one quarter?', 'MCQ', 1, 'Fractions', 'easy', 'Understand',
        json.dumps(['1/4', '1/2', '3/4', '1']), '3/4',
        'Convert one half to two quarters, then add one quarter.',
        json.dumps([{'criterion': 'Selects three quarters', 'marks': 1}]), 'page-12', 'Reviewed source excerpt'])
    sheet.freeze_panes = 'A2'
    from openpyxl.styles import Alignment, Font
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 28
        for cell in column:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if cell.row == 1:
                cell.font = Font(bold=True)
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    return output
